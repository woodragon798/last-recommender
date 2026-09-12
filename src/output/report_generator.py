"""
Output Layer
- 추천 결과를 확인하기 편한 .docx 리포트로 저장한다.
- 데이터 소스가 정상(trends)이었는지 폴백(mostPopular)이었는지 항상 리포트 상단에
  투명하게 표시한다. (실패를 숨기지 않는다는 설계 원칙)
"""
import os
from datetime import datetime, timezone, timedelta
from typing import List, Dict

from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.config import OUTPUT_DIR

KST = timezone(timedelta(hours=9))


def _add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h


def generate_report(items: List[Dict], data_source_status: str) -> str:
    today = datetime.now(KST)
    date_str = today.strftime("%Y-%m-%d")

    doc = Document()

    title = doc.add_heading(f"오늘의 사업 아이템 추천 - {date_str}", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    status_text = (
        "✅ 정상 (Google Trends 기반)"
        if data_source_status == "ok"
        else "⚠️ 폴백 모드 (YouTube 인기 차트 기반 - 트렌드 신선도가 평소보다 낮을 수 있음)"
    )
    status_para = doc.add_paragraph()
    status_run = status_para.add_run(f"데이터 소스 상태: {status_text}")
    status_run.italic = True
    status_run.font.size = Pt(10)

    disclaimer = doc.add_paragraph()
    disclaimer_run = disclaimer.add_run(
        "본 리포트는 자동 생성된 참고 자료이며, 실제 실행 여부와 최종 판단은 본인이 직접 검토 후 결정해야 합니다."
    )
    disclaimer_run.font.size = Pt(9)
    disclaimer_run.font.color.rgb = None

    doc.add_paragraph()  # spacing

    for idx, item in enumerate(items, start=1):
        _add_heading(doc, f"{idx}. {item.get('name', '(이름 없음)')}", level=1)

        meta = doc.add_paragraph()
        meta.add_run(f"카테고리: ").bold = True
        meta.add_run(f"{item.get('category', '-')}    ")
        meta.add_run(f"인지도: ").bold = True
        meta.add_run(f"{item.get('awareness_level', '-')}")

        desc = doc.add_paragraph()
        desc.add_run("설명: ").bold = True
        desc.add_run(item.get("description", "-"))

        why = doc.add_paragraph()
        why.add_run("왜 지금인가: ").bold = True
        why.add_run(item.get("why_now", "-"))

        scores = doc.add_paragraph()
        scores.add_run("초기 자본: ").bold = True
        scores.add_run(f"{item.get('capital_required', '-')}    ")
        scores.add_run("신체 활용: ").bold = True
        scores.add_run(f"{item.get('physical_involvement', '-')}")

        score_line2 = doc.add_paragraph()
        score_line2.add_run("수익 안정성 점수: ").bold = True
        score_line2.add_run(f"{item.get('income_stability_score', '-')}/5    ")
        score_line2.add_run("실행 가능성 점수: ").bold = True
        score_line2.add_run(f"{item.get('feasibility_score', '-')}/5")

        actions_heading = doc.add_paragraph()
        actions_heading.add_run("1주차 실행 계획:").bold = True
        for action in item.get("first_week_actions", []):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(action)

        doc.add_paragraph()  # 아이템 간 간격

    os.makedirs(OUTPUT_DIR, exist_ok=True)
    filepath = os.path.join(OUTPUT_DIR, f"business_recommendation_{date_str}.docx")
    doc.save(filepath)
    return filepath
